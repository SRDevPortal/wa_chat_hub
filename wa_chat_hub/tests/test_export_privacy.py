import csv
import io
import json
import unittest
from unittest.mock import patch

import frappe
from openpyxl import load_workbook
from wa_chat_hub import export_privacy as privacy


class ExportPrivacyTests(unittest.TestCase):
    def setUp(self):
        self.response_patch = patch.object(frappe.local, "response", frappe._dict())
        self.response_patch.start()
        self.addCleanup(self.response_patch.stop)

    def test_native_csv_and_excel_mask_contact_ids(self):
        # Existing fixture is only read; native exporter still checks permissions.
        name = frappe.db.get_value("Chat Contact", {}, "name")
        self.assertTrue(name)
        original = frappe.db.get_value("Chat Contact", name, "phone_number")
        for format in ("CSV", "Excel"):
            frappe.local.response = frappe._dict()
            with patch.object(privacy, "restricted", return_value=True):
                privacy.export_data("Chat Contact", all_doctypes=False, with_data=True,
                    select_columns=json.dumps({"Chat Contact": ["name", "phone_number"]}), file_type=format,
                    filters=json.dumps({"name": name}), export_without_column_meta=True)
            if format == "CSV":
                text = frappe.response["result"]
            else:
                wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
                text = str(list(wb.active.values)); wb.close()
            if original and len(original) >= 10:
                self.assertNotIn(original, text)
            self.assertEqual(frappe.db.get_value("Chat Contact", name, "phone_number"), original)

    def test_native_query_preserves_record_selection(self):
        with patch.object(frappe.local, "form_dict", frappe._dict(
            doctype="Chat Conversation", fields='["name", "contact", "status"]',
            file_format_type="CSV", selected_items='["195854"]')):
            with patch.object(privacy, "restricted", return_value=True):
                privacy.export_query()
        rows = list(csv.reader(io.StringIO(frappe.response["filecontent"].decode("utf-8-sig"))))
        self.assertEqual(len(rows), 2)
        self.assertIn("195854", rows[1])
        raw = frappe.db.get_value("Chat Conversation", "195854", "contact")
        self.assertNotIn(raw, str(rows))

    def test_unsafe_columns_rejected(self):
        for field in ("*", "raw_payload", "ai_workflow_state", "contact as raw", "tabCustomer.mobile_no"):
            with self.assertRaises(frappe.PermissionError):
                privacy.validate_fields("Chat Conversation", [field])
        privacy.validate_fields("Chat Conversation", ["`tabChat Conversation`.`contact`"])

    def test_existing_export_denial_propagates(self):
        from sriaas_clinic.api.crm_lead import privacy_outputs
        with patch.object(privacy, "restricted", return_value=True), \
             patch.object(privacy_outputs, "export_data", side_effect=frappe.PermissionError):
            with self.assertRaises(frappe.PermissionError):
                privacy.export_data("Chat Contact", select_columns={"Chat Contact": ["name"]})

    def test_full_visibility_and_other_doctypes_delegate(self):
        from sriaas_clinic.api.crm_lead import privacy_outputs
        for dt, restricted in (("Chat Contact", False), ("Customer", True)):
            with patch.object(privacy, "restricted", return_value=restricted), \
                 patch.object(privacy_outputs, "export_data", return_value="unchanged"):
                self.assertEqual(privacy.export_data(dt), "unchanged")

    def test_failure_clears_raw_download(self):
        def original():
            frappe.response.update(filecontent=b"raw", type="binary", filename="raw.xlsx")
        with self.assertRaises(frappe.ValidationError):
            privacy.protected_export(original, "Excel")
        self.assertNotIn("filecontent", frappe.response)

    def test_formula_and_embedded_numbers(self):
        result = privacy.mask_cell("=Call 9876501234")
        self.assertTrue(result.startswith("'="))
        self.assertNotIn("9876501234", result)

    def test_templates_rejected_before_export(self):
        with patch.object(privacy, "restricted", return_value=True):
            with self.assertRaises(frappe.PermissionError):
                privacy.export_data("Chat Contact", template=True)

    def test_legacy_aliases_resolve_to_same_masking_boundary(self):
        from wa_chat_hub import hooks
        for module in ("privacy_shield.outputs", "sriaas_clinic.api.crm_lead.privacy_outputs"):
            for action in ("export_data", "export_query"):
                path = module + "." + action
                target = hooks.override_whitelisted_methods[path]
                self.assertEqual(target, "wa_chat_hub.export_privacy." + action)
                endpoint = frappe.get_attr(target)
                with patch.object(privacy, "restricted", return_value=True):
                    with self.assertRaises(frappe.PermissionError):
                        if action == "export_data":
                            endpoint("Chat Conversation", select_columns={"Chat Conversation": ["raw_payload"]})
                        else:
                            with patch.object(frappe.local, "form_dict", frappe._dict(
                                doctype="Chat Conversation", fields='["raw_payload"]')):
                                endpoint()

    def test_csv_and_excel_preserve_only_exact_record_ids(self):
        from openpyxl import Workbook
        values = ["CRM-LEAD-2026-687089", "HLC-PAT-2026-38721",
                  "919876501234", "Call 919876501234 about CRM-LEAD-2026-687089",
                  "=919876501234", "CRM-LEAD-2026-687089 919876501234"]
        for file_type in ("CSV", "Excel"):
            frappe.local.response = frappe._dict()
            if file_type == "CSV":
                output = io.StringIO(); csv.writer(output).writerow(values)
                frappe.response["filecontent"] = output.getvalue().encode("utf-8-sig")
            else:
                wb = Workbook(); wb.active.append(values)
                output = io.BytesIO(); wb.save(output); wb.close()
                frappe.response["filecontent"] = output.getvalue()
            privacy.mask_download(file_type)
            if file_type == "CSV":
                row = next(csv.reader(io.StringIO(frappe.response["filecontent"].decode("utf-8-sig"))))
            else:
                wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
                row = list(next(wb.active.values)); wb.close()
            self.assertEqual(row[:2], values[:2])
            self.assertNotIn("919876501234", str(row))
            self.assertTrue(row[4].startswith("'="))

    def test_native_query_preserves_crm_reference(self):
        reference = frappe.db.get_value("Chat Conversation", "195854", "linked_reference_name")
        self.assertTrue(reference)
        for file_type in ("CSV", "Excel"):
            frappe.local.response = frappe._dict()
            with patch.object(frappe.local, "form_dict", frappe._dict(
                doctype="Chat Conversation", fields='["name", "contact", "linked_reference_name"]',
                file_format_type=file_type, selected_items='["195854"]')):
                with patch.object(privacy, "restricted", return_value=True):
                    privacy.export_query()
            if file_type == "CSV":
                text = frappe.response["filecontent"].decode("utf-8-sig")
            else:
                wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
                text = str(list(wb.active.values)); wb.close()
            self.assertIn(reference, text)
            self.assertNotIn(frappe.db.get_value("Chat Conversation", "195854", "contact"), text)
